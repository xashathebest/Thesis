"""Session-history filtering and evidence-aware CSV/XLSX export helpers."""

from __future__ import annotations

import csv
from datetime import date, datetime
from io import BytesIO, StringIO
from typing import Iterable, Mapping


# The field order is intentionally stable: it preserves the established first
# columns while adding enough source data to reproduce the final AI verdict.
EXPORT_FIELDS = {
    "fish_id": "Fish ID", "timestamp": "Timestamp", "grade": "AI Final Grade",
    "final_weighted_score": "AI Final Support (%)", "best_evidence_class": "Best Evidence Class",
    "best_evidence_score": "Best Evidence Support (%)", "verdict_status": "Verdict Status",
    "verdict_reason_code": "Verdict Reason Code", "verdict_reason": "Verdict Reason",
    "manual_grade": "Manual Grade", "manual_override": "Manual Override", "review_timestamp": "Review Timestamp",
    "detection_confidence": "Model 1 Detection Confidence (%)",
    "grade_confidence": "Accepted Final Grade Support (%)", "observed_region_count": "Observed Region Count",
    "observed_regions": "Observed Regions", "original_weight_coverage": "Original Evidence Coverage (%)",
    "evidence_completeness": "Evidence Completeness", "model2_observation_count": "Model 2 Usable Frame Observations",
    "sharpness_score": "Latest Crop Sharpness (Laplacian Variance)", "best_frame_score": "Best Frame Score",
    "best_frame_id": "Best Frame ID", "best_crop_path": "Best Crop Path",
    "body_observed": "Body Observed", "body_presence": "Body Observation Status", "body_grade": "Body Top Grade",
    "body_grade_confidence": "Body Evidence Score (%)", "body_weight": "Body Original Weight",
    "body_effective_weight": "Body Effective Weight", "body_weighted_contribution": "Body Winning Contribution (%)",
    "body_observation_count": "Body Observation Count", "body_mean_confidence": "Body Mean Evidence (%)",
    "body_max_confidence": "Body Max Evidence (%)", "body_std_dev": "Body Evidence Std Dev (%)",
    "head_observed": "Head Observed", "head_presence": "Head Observation Status", "head_grade": "Head Top Grade",
    "head_grade_confidence": "Head Evidence Score (%)", "head_weight": "Head Original Weight",
    "head_effective_weight": "Head Effective Weight", "head_weighted_contribution": "Head Winning Contribution (%)",
    "head_observation_count": "Head Observation Count", "head_mean_confidence": "Head Mean Evidence (%)",
    "head_max_confidence": "Head Max Evidence (%)", "head_std_dev": "Head Evidence Std Dev (%)",
    "tail_observed": "Tail Observed", "tail_presence": "Tail Observation Status", "tail_grade": "Tail Top Grade",
    "tail_grade_confidence": "Tail Evidence Score (%)", "tail_weight": "Tail Original Weight",
    "tail_effective_weight": "Tail Effective Weight", "tail_weighted_contribution": "Tail Winning Contribution (%)",
    "tail_observation_count": "Tail Observation Count", "tail_mean_confidence": "Tail Mean Evidence (%)",
    "tail_max_confidence": "Tail Max Evidence (%)", "tail_std_dev": "Tail Evidence Std Dev (%)",
    "class_a_weighted_score": "Class A Weighted Score (%)", "class_b_weighted_score": "Class B Weighted Score (%)",
    "class_c_weighted_score": "Class C Weighted Score (%)", "rejected_weighted_score": "Rejected Weighted Score (%)",
    "body_mean_hue": "Body Mean Hue (deg)", "body_median_hue": "Body Median Hue (deg)",
    "body_mean_saturation": "Body Mean Saturation (%)", "body_mean_value": "Body Mean Value (%)", "body_yellow_ratio": "Body Yellow Proxy (%)",
    "head_mean_hue": "Head Mean Hue (deg)", "head_median_hue": "Head Median Hue (deg)",
    "head_mean_saturation": "Head Mean Saturation (%)", "head_mean_value": "Head Mean Value (%)", "head_yellow_ratio": "Head Yellow Proxy (%)",
    "tail_mean_hue": "Tail Mean Hue (deg)", "tail_median_hue": "Tail Median Hue (deg)",
    "tail_mean_saturation": "Tail Mean Saturation (%)", "tail_mean_value": "Tail Mean Value (%)", "tail_yellow_ratio": "Tail Yellow Proxy (%)",
    "whole_fish_mean_hue": "Whole Fish Mean Hue (deg)", "whole_fish_median_hue": "Whole Fish Median Hue (deg)",
    "whole_fish_mean_saturation": "Whole Fish Mean Saturation (%)", "whole_fish_mean_value": "Whole Fish Mean Value (%)", "whole_fish_yellow_ratio": "Whole Fish Yellow Proxy (%)",
    "detected_regions": "Raw Model 2 Detected Regions", "model2_classes": "Raw Model 2 Classes",
    "adjustment_applied": "Adjustment Applied", "adjustment_reason": "Adjustment Reason",
    "override_applied": "Rejected Override Applied", "override_reason": "Rejected Override Reason",
    "model1_version": "Model 1 Checkpoint", "model1_hash": "Model 1 SHA256 (short)",
    "model2_version": "Model 2 Checkpoint", "model2_hash": "Model 2 SHA256 (short)",
    "grading_config_version": "Grading Config Version", "detection_threshold": "Detection Threshold",
    "quality_threshold": "Quality Evidence Threshold", "final_verdict_threshold": "Final Verdict Threshold",
    "minimum_coverage_setting": "Minimum Coverage Setting", "grading_mode": "Grading Mode",
    "final_explanation": "Calculation / Explanation", "processing_time": "Processing Time (ms)",
    # Retain every aggregated Model 2 region/grade support value so a CSV/XLSX
    # export can later be replayed with different weights or acceptance rules.
    # These are evidence/support scores, not calibrated class probabilities.
    "body_presence_confidence": "Body Presence Evidence (%)",
    "body_class_a_evidence": "Body Class A Evidence (%)", "body_class_b_evidence": "Body Class B Evidence (%)",
    "body_class_c_evidence": "Body Class C Evidence (%)", "body_rejected_evidence": "Body Rejected Evidence (%)",
    "head_presence_confidence": "Head Presence Evidence (%)",
    "head_class_a_evidence": "Head Class A Evidence (%)", "head_class_b_evidence": "Head Class B Evidence (%)",
    "head_class_c_evidence": "Head Class C Evidence (%)", "head_rejected_evidence": "Head Rejected Evidence (%)",
    "tail_presence_confidence": "Tail Presence Evidence (%)",
    "tail_class_a_evidence": "Tail Class A Evidence (%)", "tail_class_b_evidence": "Tail Class B Evidence (%)",
    "tail_class_c_evidence": "Tail Class C Evidence (%)", "tail_rejected_evidence": "Tail Rejected Evidence (%)",
}
DEFAULT_EXPORT_FIELDS = tuple(EXPORT_FIELDS)


def parse_timestamp(value: object) -> datetime | None:
    try:
        return datetime.fromisoformat(str(value))
    except (TypeError, ValueError):
        return None


def _number(value: object, *, scale: float = 1.0) -> float | str:
    try:
        number = float(value) * scale
    except (TypeError, ValueError):
        return ""
    return round(number, 4) if number == number and abs(number) != float("inf") else ""


def _part_values(event: Mapping[str, object], key: str) -> str:
    parts = event.get("parts")
    if not isinstance(parts, list):
        return ""
    return ", ".join(sorted({str(part.get(key, "")) for part in parts if isinstance(part, dict) and part.get(key)}))


def _analysis(event: Mapping[str, object]) -> dict[str, object]:
    value = event.get("analysis")
    return dict(value) if isinstance(value, dict) else {}


def _part_analysis(analysis: Mapping[str, object], region: str) -> dict[str, object]:
    parts = analysis.get("part_results")
    value = parts.get(region) if isinstance(parts, dict) else None
    return dict(value) if isinstance(value, dict) else {}


def _color_analysis(analysis: Mapping[str, object], region: str) -> dict[str, object]:
    colors = analysis.get("color")
    value = colors.get(region) if isinstance(colors, dict) else None
    return dict(value) if isinstance(value, dict) else {}


def _top_statistics(part: Mapping[str, object]) -> dict[str, object]:
    temporal = part.get("temporal_statistics")
    if not isinstance(temporal, dict):
        return {}
    top = temporal.get("top_grade_statistics")
    return dict(top) if isinstance(top, dict) else {}


def _traceability(analysis: Mapping[str, object]) -> dict[str, object]:
    trace = analysis.get("traceability")
    return dict(trace) if isinstance(trace, dict) else {}


def inspection_row(event: dict[str, object]) -> dict[str, object]:
    """Map one completed fish to numeric, reconstructable export values."""

    analysis = _analysis(event)
    weighted = analysis.get("weighted_scores") if isinstance(analysis.get("weighted_scores"), dict) else {}
    final_grade = event.get("quality") or analysis.get("final_grade") or "Ungraded"
    best_grade = analysis.get("provisional_grade", analysis.get("best_evidence_class"))
    best_score = analysis.get("provisional_score", analysis.get("best_evidence_score"))
    final_score = analysis.get("final_score")
    model1 = analysis.get("model1_detection") if isinstance(analysis.get("model1_detection"), dict) else {}
    trace = _traceability(analysis)
    frame_quality = analysis.get("frame_quality") if isinstance(analysis.get("frame_quality"), dict) else {}
    latest_quality = frame_quality.get("latest") if isinstance(frame_quality.get("latest"), dict) else {}
    best_frame = analysis.get("best_frame") if isinstance(analysis.get("best_frame"), dict) else {}
    row: dict[str, object] = {
        "fish_id": event.get("fish_label", f"Fish #{event.get('track_id', '')}"),
        "timestamp": event.get("timestamp", ""),
        "grade": final_grade,
        # Support remains visible for ungraded results as "best evidence";
        # it is not mislabeled as an accepted grade confidence.
        "final_weighted_score": _number(final_score if final_score is not None else best_score, scale=100),
        "best_evidence_class": best_grade or "",
        "best_evidence_score": _number(best_score, scale=100),
        "verdict_status": analysis.get("verdict_status", ""),
        "verdict_reason_code": analysis.get("verdict_reason_code", ""),
        "verdict_reason": analysis.get("verdict_reason_text", ""),
        "manual_grade": event.get("manual_grade") or "",
        "manual_override": "Yes" if event.get("manual_override") else "No",
        "review_timestamp": event.get("review_timestamp") or "",
        "detection_confidence": _number(model1.get("confidence") if isinstance(model1, dict) else event.get("final_confidence"), scale=100),
        "grade_confidence": _number(event.get("quality_confidence"), scale=100),
        "observed_region_count": len(analysis.get("observed_regions", ())) if isinstance(analysis.get("observed_regions"), list) else "",
        "observed_regions": ", ".join(str(region) for region in analysis.get("observed_regions", ()) if region),
        "original_weight_coverage": _number(analysis.get("original_weight_coverage"), scale=100),
        "evidence_completeness": analysis.get("evidence_completeness", ""),
        "model2_observation_count": analysis.get("observation_count", ""),
        "sharpness_score": _number(latest_quality.get("sharpness_score") if isinstance(latest_quality, dict) else None),
        "best_frame_score": _number(best_frame.get("best_frame_score") if isinstance(best_frame, dict) else None),
        "best_frame_id": best_frame.get("best_frame_id", "") if isinstance(best_frame, dict) else "",
        "best_crop_path": best_frame.get("best_crop_path", "") if isinstance(best_frame, dict) else "",
        "class_a_weighted_score": _number(weighted.get("Class A") if isinstance(weighted, dict) else None, scale=100),
        "class_b_weighted_score": _number(weighted.get("Class B") if isinstance(weighted, dict) else None, scale=100),
        "class_c_weighted_score": _number(weighted.get("Class C") if isinstance(weighted, dict) else None, scale=100),
        "rejected_weighted_score": _number(weighted.get("Rejected") if isinstance(weighted, dict) else None, scale=100),
        "detected_regions": _part_values(event, "region"),
        "model2_classes": _part_values(event, "source_class_name"),
        "adjustment_applied": "Yes" if analysis.get("adjustments") else "No",
        "adjustment_reason": "; ".join(str(item.get("reason", "")) for item in analysis.get("adjustments", ()) if isinstance(item, dict)),
        "override_applied": "Yes" if isinstance(analysis.get("override"), dict) and analysis["override"].get("applied") else "No",
        "override_reason": analysis.get("override", {}).get("reason", "") if isinstance(analysis.get("override"), dict) else "",
        "model1_version": trace.get("model1_checkpoint", ""),
        "model1_hash": trace.get("model1_checkpoint_sha256", ""),
        "model2_version": trace.get("model2_checkpoint", ""),
        "model2_hash": trace.get("model2_checkpoint_sha256", ""),
        "grading_config_version": trace.get("grading_config_version", ""),
        "detection_threshold": _number(trace.get("detection_threshold")),
        "quality_threshold": _number(trace.get("quality_threshold")),
        "final_verdict_threshold": _number(trace.get("final_verdict_threshold")),
        "minimum_coverage_setting": _number(trace.get("minimum_original_weight_coverage")),
        "grading_mode": trace.get("grading_mode", ""),
        "final_explanation": " ".join(str(item) for item in analysis.get("explanation", ()) if item),
        "processing_time": event.get("processing_time_ms") if event.get("processing_time_ms") is not None else "",
    }
    winning_grade = str(final_grade if final_grade != "Ungraded" else (best_grade or ""))
    for region, prefix in (("Body", "body"), ("Head", "head"), ("Tail", "tail")):
        part, color, stats = _part_analysis(analysis, region), _color_analysis(analysis, region), _top_statistics(_part_analysis(analysis, region))
        contributions = part.get("contributions") if isinstance(part.get("contributions"), dict) else {}
        grade_evidence = part.get("grade_evidence") if isinstance(part.get("grade_evidence"), dict) else {}
        row[f"{prefix}_observed"] = "Yes" if part.get("present") is True else "No"
        row[f"{prefix}_presence"] = part.get("status", part.get("observation_status", part.get("missing_status", "unknown")))
        row[f"{prefix}_presence_confidence"] = _number(part.get("presence_confidence"), scale=100)
        row[f"{prefix}_class_a_evidence"] = _number(grade_evidence.get("Class A"), scale=100)
        row[f"{prefix}_class_b_evidence"] = _number(grade_evidence.get("Class B"), scale=100)
        row[f"{prefix}_class_c_evidence"] = _number(grade_evidence.get("Class C"), scale=100)
        row[f"{prefix}_rejected_evidence"] = _number(grade_evidence.get("Rejected"), scale=100)
        row[f"{prefix}_grade"] = part.get("grade", "")
        row[f"{prefix}_grade_confidence"] = _number(part.get("grade_confidence"), scale=100)
        row[f"{prefix}_weight"] = _number(part.get("original_weight", part.get("weight")))
        row[f"{prefix}_effective_weight"] = _number(part.get("effective_weight", part.get("normalized_weight")))
        row[f"{prefix}_weighted_contribution"] = _number(contributions.get(winning_grade) if isinstance(contributions, dict) else None, scale=100)
        row[f"{prefix}_observation_count"] = stats.get("number_of_valid_observations", "")
        row[f"{prefix}_mean_confidence"] = _number(stats.get("mean_confidence"), scale=100)
        row[f"{prefix}_max_confidence"] = _number(stats.get("max_confidence"), scale=100)
        row[f"{prefix}_std_dev"] = _number(stats.get("standard_deviation"), scale=100)
        row[f"{prefix}_mean_hue"] = _number(color.get("mean_hue_deg"))
        row[f"{prefix}_median_hue"] = _number(color.get("median_hue_deg"))
        row[f"{prefix}_mean_saturation"] = _number(color.get("mean_saturation"), scale=100)
        row[f"{prefix}_mean_value"] = _number(color.get("mean_value"), scale=100)
        row[f"{prefix}_yellow_ratio"] = _number(color.get("yellow_ratio_proxy"), scale=100)
    whole = _color_analysis(analysis, "Whole Fish")
    row["whole_fish_mean_hue"] = _number(whole.get("mean_hue_deg"))
    row["whole_fish_median_hue"] = _number(whole.get("median_hue_deg"))
    row["whole_fish_mean_saturation"] = _number(whole.get("mean_saturation"), scale=100)
    row["whole_fish_mean_value"] = _number(whole.get("mean_value"), scale=100)
    row["whole_fish_yellow_ratio"] = _number(whole.get("yellow_ratio_proxy"), scale=100)
    return row


def filter_events(events: Iterable[dict[str, object]], *, range_name: str = "current_session", start_date: str | None = None, end_date: str | None = None) -> list[dict[str, object]]:
    values = list(events)
    if range_name == "current_session":
        return values
    if range_name == "today":
        today = date.today()
        return [event for event in values if (parsed := parse_timestamp(event.get("timestamp"))) and parsed.date() == today]
    if range_name != "custom":
        raise ValueError("Unsupported export range.")
    try:
        lower = date.fromisoformat(start_date) if start_date else None
        upper = date.fromisoformat(end_date) if end_date else None
    except ValueError as exc:
        raise ValueError("Custom export dates must use YYYY-MM-DD.") from exc
    if lower and upper and lower > upper:
        raise ValueError("The start date must not be after the end date.")
    return [
        event for event in values
        if (parsed := parse_timestamp(event.get("timestamp"))) is not None
        and (not lower or parsed.date() >= lower) and (not upper or parsed.date() <= upper)
    ]


def make_csv(events: Iterable[dict[str, object]], fields: Iterable[str]) -> bytes:
    selected = tuple(fields)
    buffer = StringIO(newline="")
    writer = csv.DictWriter(buffer, fieldnames=[EXPORT_FIELDS[field] for field in selected], extrasaction="ignore")
    writer.writeheader()
    for event in events:
        row = inspection_row(event)
        writer.writerow({EXPORT_FIELDS[field]: row[field] for field in selected})
    return buffer.getvalue().encode("utf-8-sig")


def _style_header(sheet) -> None:
    from openpyxl.styles import Font, PatternFill
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")


def make_xlsx(events: Iterable[dict[str, object]], fields: Iterable[str], summary: dict[str, object], grading_rules: Mapping[str, object] | None = None) -> bytes:
    """Create all five explainable inspection workbook sheets."""

    try:
        from openpyxl import Workbook
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("Excel export requires openpyxl. Install the project requirements.") from exc
    selected, event_rows = tuple(fields), list(events)
    workbook = Workbook()
    inspections = workbook.active
    inspections.title = "Fish Inspections"
    inspections.append([EXPORT_FIELDS[field] for field in selected])
    _style_header(inspections)
    for event in event_rows:
        row = inspection_row(event)
        inspections.append([row[field] for field in selected])
    inspections.freeze_panes = "A2"
    for column in inspections.columns:
        inspections.column_dimensions[column[0].column_letter].width = min(48, max(14, max(len(str(cell.value or "")) for cell in column) + 2))

    summary_sheet = workbook.create_sheet("Session Summary")
    summary_sheet.append(["Metric", "Value"])
    _style_header(summary_sheet)
    for label, value in summary.items():
        summary_sheet.append([label, value])
    summary_sheet.column_dimensions["A"].width, summary_sheet.column_dimensions["B"].width = 42, 26

    feature_sheet = workbook.create_sheet("Feature Summary")
    feature_sheet.append(["Feature", "Count", "Percentage of fish", "Source"])
    _style_header(feature_sheet)
    total = len(event_rows)
    for region in ("Head", "Body", "Tail"):
        count = sum(int(_part_analysis(_analysis(event), region).get("present") is True) for event in event_rows)
        feature_sheet.append([f"{region} confidently observed", count, round(count / total * 100, 2) if total else 0.0, "Model 2 part detector"])
    feature_sheet.append(["HSV colour metrics", total, 100.0 if total else 0.0, "Descriptive proxy; no active grading rule"])
    feature_sheet.append(["Trained defect labels", 0, 0.0, "None in supplied Model 1/Model 2 checkpoints"])
    for column, width in {"A": 30, "B": 14, "C": 22, "D": 56}.items():
        feature_sheet.column_dimensions[column].width = width

    rules = dict(grading_rules or {})
    weights = rules.get("part_weights") if isinstance(rules.get("part_weights"), dict) else {}
    rule_sheet = workbook.create_sheet("Grade Calculation Rules")
    rule_sheet.append(["Rule", "Value", "Status"])
    _style_header(rule_sheet)
    for region in ("Body", "Head", "Tail"):
        rule_sheet.append([f"{region} original weight", weights.get(region, ""), "Active"])
    rule_sheet.append(["Aggregation", "Strongest Model 2 confidence per region/grade per usable frame; mean over observed frames", "Active"])
    rule_sheet.append(["Unavailable region", "Exclude and renormalize observed region weights; do not mark physically missing", "Active"])
    rule_sheet.append(["Body required", rules.get("require_body", True), "Active"])
    rule_sheet.append(["Grading mode", rules.get("grading_mode", "standard"), "Active"])
    rule_sheet.append(["Minimum regions", rules.get("minimum_regions_observed", ""), "Active"])
    rule_sheet.append(["Minimum original coverage", rules.get("minimum_original_weight_coverage", ""), "Active"])
    rule_sheet.append(["Final verdict threshold", rules.get("active_final_verdict_threshold", rules.get("final_verdict_threshold", "")), "Configurable"])
    rule_sheet.append(["Minimum valid frames", rules.get("minimum_track_observations", ""), "Active"])
    rule_sheet.append(["Rejected override threshold", rules.get("rejected_override_threshold", "None"), "Disabled" if rules.get("rejected_override_threshold") is None else "Configured"])
    rule_sheet.append(["HSV adjustments", rules.get("color_adjustments_enabled", False), "No validated rule configured"])
    rule_sheet.append(["Defect adjustments", rules.get("defect_adjustments_enabled", False), "No trained defect label or validated rule configured"])
    rule_sheet.column_dimensions["A"].width, rule_sheet.column_dimensions["B"].width, rule_sheet.column_dimensions["C"].width = 30, 96, 42

    review_sheet = workbook.create_sheet("Review Queue")
    review_headers = ["Fish ID", "Timestamp", "Best Evidence Class", "Best Score (%)", "Evidence Coverage (%)", "Observed Regions", "Review Reason", "Manual Grade"]
    review_sheet.append(review_headers)
    _style_header(review_sheet)
    for event in event_rows:
        analysis = _analysis(event)
        if analysis.get("verdict_status") != "NEEDS_REVIEW":
            continue
        row = inspection_row(event)
        review_sheet.append([
            row["fish_id"], row["timestamp"], row["best_evidence_class"], row["best_evidence_score"],
            row["original_weight_coverage"], row["observed_regions"], row["verdict_reason"], row["manual_grade"],
        ])
    for column, width in {"A": 14, "B": 27, "C": 22, "D": 16, "E": 22, "F": 28, "G": 74, "H": 18}.items():
        review_sheet.column_dimensions[column].width = width
    review_sheet.freeze_panes = "A2"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
