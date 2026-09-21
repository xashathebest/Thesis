from __future__ import annotations

from datetime import datetime, timedelta
import tempfile
import unittest
from pathlib import Path

from src.api.database import InspectionDatabase, InspectionSessionContext, grading_config_hash
from src.api.domain import Detection, QualitySummary, TrackingConfig, TrackingManager
from src.api.export import DEFAULT_EXPORT_FIELDS, make_csv, make_xlsx


def event(
    track_id: int,
    *,
    timestamp: str | None = None,
    quality: str = "Ungraded",
    manual_grade: str | None = None,
) -> dict[str, object]:
    final_score = 0.78 if quality == "Class A" else None
    return {
        "track_id": track_id,
        "fish_label": f"Fish #{track_id}",
        "timestamp": timestamp or datetime.now().astimezone().isoformat(timespec="seconds"),
        "final_class": "Fish",
        "final_confidence": 0.91,
        "quality": quality,
        "quality_confidence": final_score,
        "parts": [],
        "manual_override": manual_grade is not None,
        "manual_grade": manual_grade,
        "analysis": {
            "final_grade": quality,
            "final_score": final_score,
            "second_grade": "Class B",
            "second_support": 0.21,
            "grade_margin": 0.57 if final_score is not None else None,
            "original_weight_coverage": 1.0,
            "observed_regions": ["Body", "Head", "Tail"],
            "verdict_reason_code": "GR_CONFIDENT" if quality != "Ungraded" else "UG_LOW_GRADE_SUPPORT",
            "weighted_scores": {"Class A": 0.78, "Class B": 0.21, "Class C": 0.0, "Rejected": 0.0},
            "part_results": {"Body": {"present": True}},
            "color": {"Body": {"status": "measured"}},
            "model1_detection": {"confidence": 0.91},
        },
        "processing_time_ms": 11.4,
    }


class InspectionDatabaseTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.path = Path(self.temporary_directory.name) / "history.sqlite3"
        self.database = InspectionDatabase(self.path)

    def tearDown(self) -> None:
        self.temporary_directory.cleanup()

    @staticmethod
    def metadata() -> dict[str, object]:
        grading = {"part_weights": {"Body": 0.5, "Head": 0.3, "Tail": 0.2}, "final_verdict_threshold": 0.5}
        return {
            "git_commit": "abc123",
            "model1_path": "models/fish.pt",
            "model1_sha256": "model-one-hash",
            "model2_path": "models/quality.pt",
            "model2_sha256": "model-two-hash",
            "tracker_backend": "ByteTrack",
            "grading_config": grading,
            "grading_config_hash": grading_config_hash(grading),
            "camera_settings": {"available": True, "actual": {"exposure": -6}},
            "runtime_settings": {"quality_interval": 3, "roi_padding": 0},
        }

    def test_event_upsert_survives_reopening_and_keeps_session_snapshot(self) -> None:
        session_id = self.database.create_session(self.metadata())
        self.database.upsert_event(session_id, event(7))
        self.database.upsert_event(session_id, event(7, quality="Class A", manual_grade="Class A"))

        reopened = InspectionDatabase(self.path)
        records = reopened.events(session_id=session_id, range_name="current_session")
        self.assertEqual(len(records), 1)
        self.assertEqual(records[0]["track_id"], 7)
        self.assertEqual(records[0]["quality"], "Class A")
        self.assertEqual(records[0]["manual_grade"], "Class A")
        self.assertEqual(records[0]["session_id"], session_id)

        saved_session = reopened.session(session_id)
        self.assertIsNotNone(saved_session)
        assert saved_session is not None
        self.assertEqual(saved_session["git_commit"], "abc123")
        self.assertEqual(saved_session["tracker_backend"], "ByteTrack")
        self.assertEqual(saved_session["camera_settings"], {"actual": {"exposure": -6}, "available": True})

    def test_durable_today_custom_and_session_filters(self) -> None:
        first = self.database.create_session(self.metadata())
        second = self.database.create_session(self.metadata())
        now = datetime.now().astimezone()
        old = now - timedelta(days=3)
        self.database.upsert_event(first, event(1, timestamp=old.isoformat(timespec="seconds")))
        self.database.upsert_event(second, event(2, timestamp=now.isoformat(timespec="seconds"), quality="Class A"))

        self.assertEqual([item["track_id"] for item in self.database.events(session_id=first, range_name="current_session")], [1])
        self.assertEqual([item["track_id"] for item in self.database.events(range_name="today")], [2])
        self.assertEqual(
            [item["track_id"] for item in self.database.events(range_name="custom", start_date=old.date().isoformat(), end_date=old.date().isoformat())],
            [1],
        )

    def test_tracking_listener_persists_created_late_grade_and_review(self) -> None:
        session_id = self.database.create_session(self.metadata())
        tracking = TrackingManager(
            TrackingConfig(line_orientation="vertical", line_position=0.5, conveyor_direction="left_to_right")
        )
        tracking.add_event_listener(lambda _action, payload: self.database.upsert_event(session_id, payload))
        left = Detection(0, "Fish", 0.9, (5, 10, 25, 35), track_id=13)
        right = Detection(0, "Fish", 0.9, (65, 10, 85, 35), track_id=13)
        tracking.update([left], (100, 60), timestamp=0.0)
        tracking.update([right], (100, 60), timestamp=1.0, event_timestamp="2026-01-01T12:00:00+00:00")
        grade = QualitySummary(
            13,
            "Class A",
            0.78,
            analysis={"final_grade": "Class A", "final_score": 0.78, "verdict_reason_code": "GR_CONFIDENT"},
        )
        tracking.update([right], (100, 60), timestamp=1.1, grades={13: grade})
        tracking.apply_manual_review(13, "Class A")

        records = self.database.events(session_id=session_id, range_name="current_session")
        self.assertEqual(len(records), 1)
        self.assertEqual(records[0]["quality"], "Class A")
        self.assertEqual(records[0]["manual_grade"], "Class A")

    def test_session_context_closes_old_session_when_rotated(self) -> None:
        context = InspectionSessionContext(self.database)
        first = context.start(self.metadata())
        second = context.rotate(self.metadata())
        self.assertNotEqual(first, second)
        first_session = self.database.session(first)
        self.assertIsNotNone(first_session)
        assert first_session is not None
        self.assertIsNotNone(first_session["ended_at"])
        self.assertEqual(context.event_session(self.metadata()), second)

    def test_csv_and_xlsx_use_the_same_durable_event_and_session_snapshot(self) -> None:
        session_id = self.database.create_session(self.metadata())
        self.database.upsert_event(session_id, event(19, quality="Class A"))
        records = self.database.events(session_id=session_id, range_name="current_session")
        csv_payload = make_csv(records, DEFAULT_EXPORT_FIELDS).decode("utf-8-sig")
        self.assertIn("Session ID", csv_payload)
        self.assertIn(session_id, csv_payload)

        workbook = make_xlsx(
            records,
            DEFAULT_EXPORT_FIELDS,
            {"Total Fish": 1},
            session_configuration=self.database.session(session_id),
        )
        from io import BytesIO
        from openpyxl import load_workbook

        loaded = load_workbook(BytesIO(workbook))
        self.assertIn("Session Configuration", loaded.sheetnames)
        self.assertEqual(loaded["Session Configuration"].cell(row=2, column=1).value, "session_id")


if __name__ == "__main__":
    unittest.main()
