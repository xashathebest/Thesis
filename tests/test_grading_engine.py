"""Unit coverage for explainable, body-weighted fish grading."""

from __future__ import annotations

import csv
import json
from io import BytesIO, StringIO
import unittest

import numpy as np

from src.api.export import DEFAULT_EXPORT_FIELDS, inspection_row, make_csv, make_xlsx
from src.api.domain import Detection, QualitySummary, TrackingConfig, TrackingManager
from src.inference.grading_engine import GradingConfig, WeightedGradingEngine
from src.inference.part_types import PartDetection


CLASS_IDS = {
    ("Body", "Class A"): 0, ("Head", "Class A"): 1, ("Tail", "Class A"): 2,
    ("Body", "Class B"): 3, ("Head", "Class B"): 4, ("Tail", "Class B"): 5,
    ("Body", "Class C"): 6, ("Head", "Class C"): 7, ("Tail", "Class C"): 8,
    ("Body", "Rejected"): 9, ("Head", "Rejected"): 10, ("Tail", "Rejected"): 11,
}
BOXES = {"Head": (2, 8, 18, 30), "Body": (18, 6, 46, 33), "Tail": (46, 9, 60, 28)}


def part(region: str, grade: str, confidence: float) -> PartDetection:
    return PartDetection.from_source_class(CLASS_IDS[(region, grade)], confidence, BOXES[region])


def crop() -> np.ndarray:
    image = np.zeros((40, 64, 3), dtype=np.uint8)
    image[:, :, :] = (70, 145, 190)  # BGR; real pixels used only for HSV measurement.
    return image


class GradingEngineTests(unittest.TestCase):
    def test_documented_nested_configuration_is_equivalent_to_the_flat_rules(self) -> None:
        config = GradingConfig.from_mapping(
            {
                "config_version": "nested-test",
                "grading": {
                    "weights": {"body": 0.50, "head": 0.30, "tail": 0.20},
                    "require_body": True,
                    "minimum_regions_observed": 1,
                    "minimum_original_weight_coverage": 0.50,
                    "final_verdict_threshold": 0.50,
                    "temporal": {"minimum_valid_frames": 2, "use_frame_quality_filter": True},
                    "frame_quality": {"minimum_crop_area": 48},
                    "best_frame": {"best_frame_sharpness_weight": 0.40},
                    "rejected_override": {"enabled": False, "threshold": 0.85},
                    "hsv_adjustment": {"enabled": False},
                },
                "storage": {"save_best_fish_crop": True, "best_crop_directory": "results/crops"},
            }
        )
        self.assertEqual(config.config_version, "nested-test")
        self.assertEqual(config.minimum_track_observations, 2)
        self.assertTrue(config.use_frame_quality_filter)
        self.assertEqual(config.minimum_crop_area, 48)
        self.assertAlmostEqual(config.best_frame_sharpness_weight, 0.40)
        self.assertIsNone(config.rejected_override_threshold)
        self.assertFalse(config.color_adjustments_enabled)
        self.assertTrue(config.save_best_fish_crop)
        self.assertEqual(config.best_crop_directory, "results/crops")

    def test_meaningful_unknown_or_conflicting_legacy_threshold_is_rejected(self) -> None:
        with self.assertRaises(ValueError):
            GradingConfig.from_mapping({"grading": {"final_verdict_threshold": .50, "final_verdict_threhsold": .40}})
        with self.assertRaises(ValueError):
            GradingConfig.from_mapping({"grading": {"final_verdict_threshold": .50, "minimum_final_score": .25}})

    def engine(self, **overrides: object) -> WeightedGradingEngine:
        return WeightedGradingEngine(GradingConfig(**overrides))

    def test_documented_all_part_formula_and_contributions(self) -> None:
        evidence = [
            part("Body", "Class A", .59), part("Body", "Class B", .25), part("Body", "Class C", .10), part("Body", "Rejected", .06),
            part("Head", "Class A", .72), part("Head", "Class B", .18), part("Head", "Class C", .07), part("Head", "Rejected", .03),
            part("Tail", "Class A", .64), part("Tail", "Class B", .22), part("Tail", "Class C", .10), part("Tail", "Rejected", .04),
        ]
        verdict = self.engine().evaluate(128, crop(), evidence, stabilize=False)
        self.assertEqual(verdict.final_grade, "Class A")
        self.assertAlmostEqual(verdict.weighted_scores["Class A"], .639)
        self.assertAlmostEqual(verdict.weighted_scores["Class B"], .223)
        self.assertAlmostEqual(verdict.part_results["Body"]["contributions"]["Class A"], .295)
        self.assertAlmostEqual(verdict.part_results["Head"]["contributions"]["Class A"], .216)
        self.assertAlmostEqual(verdict.part_results["Tail"]["contributions"]["Class A"], .128)

    def test_mixed_grades_compare_all_four_weighted_scores(self) -> None:
        verdict = self.engine(final_verdict_threshold=.40).evaluate(9, crop(), [
            part("Body", "Class B", .80), part("Head", "Class A", .90), part("Tail", "Class A", .90),
        ], stabilize=False)
        self.assertAlmostEqual(verdict.weighted_scores["Class A"], .45)
        self.assertAlmostEqual(verdict.weighted_scores["Class B"], .40)
        self.assertEqual(verdict.final_grade, "Class A")

    def test_unknown_tail_keeps_fixed_weights_and_is_not_marked_physically_missing(self) -> None:
        verdict = self.engine().evaluate(10, crop(), [part("Body", "Class A", .60), part("Head", "Class A", .80)], stabilize=False)
        self.assertEqual(verdict.evidence_completeness, "partial")
        self.assertAlmostEqual(verdict.part_results["Body"]["normalized_weight"], .50)
        self.assertAlmostEqual(verdict.part_results["Head"]["normalized_weight"], .30)
        self.assertEqual(verdict.part_results["Tail"]["missing_status"], "unknown_not_observed")
        self.assertIsNone(verdict.part_results["Tail"]["physical_missing"])
        self.assertAlmostEqual(verdict.weighted_scores["Class A"], .54)

    def test_low_confidence_or_no_body_returns_ungraded(self) -> None:
        low = self.engine().evaluate(1, crop(), [part("Body", "Class A", .20), part("Head", "Class A", .95)], stabilize=False)
        self.assertIsNone(low.final_grade)
        missing_body = self.engine().evaluate(2, crop(), [part("Head", "Class A", .95), part("Tail", "Class A", .95)], stabilize=False)
        self.assertIsNone(missing_body.final_grade)

    def test_configured_rejected_label_override_is_explicit(self) -> None:
        engine = self.engine(rejected_override_threshold=.90)
        verdict = engine.evaluate(11, crop(), [
            part("Body", "Class A", .98), part("Head", "Class A", .98), part("Tail", "Class A", .98),
            part("Body", "Rejected", .94),
        ], stabilize=False)
        self.assertEqual(verdict.final_grade, "Rejected")
        self.assertTrue(verdict.override and verdict.override["applied"])
        self.assertEqual(verdict.override["source"], "Model 2 trained Rejected_* part label")

    def test_track_evidence_is_mean_of_nonduplicated_frame_scores(self) -> None:
        engine = self.engine(minimum_track_observations=2, final_verdict_threshold=.25)
        first = engine.evaluate(44, crop(), [part("Body", "Class A", .55)], stabilize=True)
        second = engine.evaluate(44, crop(), [part("Body", "Class A", .65)], stabilize=True)
        self.assertIsNone(first.final_grade)
        self.assertEqual(second.final_grade, "Class A")
        self.assertAlmostEqual(second.final_score or 0, .30)
        # Another fish's evidence has its own Model 1 ID and cannot mix here.
        other = engine.evaluate(45, crop(), [part("Body", "Class B", .92)], stabilize=False)
        self.assertEqual(other.final_grade, "Class B")

    def test_color_is_measured_but_has_no_unconfigured_grade_adjustment(self) -> None:
        verdict = self.engine().evaluate(3, crop(), [part("Body", "Class C", .70)], stabilize=False)
        body_color = verdict.color["Body"]
        self.assertEqual(body_color["status"], "measured")
        self.assertEqual(body_color["grade_adjustment"], 0.0)
        self.assertEqual(verdict.adjustments, ())

    def test_supporting_measurements_are_optional_traceable_and_do_not_change_grades(self) -> None:
        evidence = [part("Body", "Class A", .9), part("Head", "Class A", .9)]
        enabled = self.engine(image_diagnostics_enabled=True)
        disabled = self.engine(image_diagnostics_enabled=False)
        for frame_id in (10, 11):
            on = enabled.evaluate(1, crop(), evidence, stabilize=True, frame_id=frame_id)
            off = disabled.evaluate(1, crop(), evidence, stabilize=True, frame_id=frame_id)
        self.assertEqual(on.final_grade, off.final_grade)
        self.assertEqual(on.weighted_scores, off.weighted_scores)
        self.assertEqual(on.reason_codes, off.reason_codes)
        payload = on.to_dict()["supporting_image_measurements"]
        self.assertEqual(payload["fish_id"], 1)
        self.assertFalse(payload["used_for_final_grade"])
        self.assertEqual([item["frame_id"] for item in payload["regions"]["Body"]["frame_rois"]], [10, 11])
        self.assertIn("no segmentation mask", payload["regions"]["Body"]["measurement_source"])
        self.assertEqual(payload["texture"]["missing_value_reason"], "no_validated_fish_mask_in_runtime")
        self.assertEqual(off.supporting_image_measurements["regions"]["Body"]["missing_value_reason"], "diagnostics_disabled")
        json.dumps(payload, allow_nan=False)
        other = enabled.evaluate(2, np.full((40, 64, 3), 255, dtype=np.uint8), evidence, stabilize=False, frame_id="other")
        other_payload = other.supporting_image_measurements
        self.assertEqual(other_payload["fish_id"], 2)
        self.assertEqual(other_payload["regions"]["Body"]["frame_rois"][0]["frame_id"], "other")
        self.assertNotEqual(payload["regions"]["Body"]["values"]["mean_value"], other_payload["regions"]["Body"]["values"]["mean_value"])
        row = inspection_row({"track_id": 1, "quality": on.final_grade, "analysis": on.to_dict()})
        self.assertEqual(json.loads(row["supporting_image_measurements"])["fish_id"], 1)

    def test_invalid_diagnostic_crop_is_unavailable_without_inventing_pixels(self) -> None:
        result = self.engine().evaluate(3, np.empty((0, 0, 3), dtype=np.uint8), [part("Body", "Class A", .9)], stabilize=False, frame_id="invalid")
        whole = result.supporting_image_measurements["regions"]["Whole Fish"]
        self.assertEqual(whole["status"], "unavailable")
        self.assertIn("Invalid", whole["missing_value_reason"])
        self.assertEqual(whole["frame_rois"], [])

    def test_temporal_gate_requires_two_distinct_evidence_bearing_frames(self) -> None:
        cases = (
            ("body", [part("Body", "Class A", .9)]),
            ("body_head", [part("Body", "Class A", .9), part("Head", "Class A", .9)]),
            ("body_tail", [part("Body", "Class A", .9), part("Tail", "Class A", .9)]),
        )
        for name, evidence in cases:
            with self.subTest(case=name):
                live = self.engine(minimum_track_observations=2, final_verdict_threshold=.25)
                live.evaluate(1, crop(), evidence, stabilize=True, frame_id=1)
                empty = live.evaluate(1, crop(), [], stabilize=True, frame_id=2)
                self.assertIsNone(empty.final_grade)
                self.assertIn("UG_TEMPORAL_EVIDENCE", empty.reason_codes)
                self.assertEqual(empty.temporal_stability["evidence_bearing_frame_count"], 1)
                repeated = live.evaluate(1, crop(), evidence, stabilize=True, frame_id=1)
                self.assertIsNone(repeated.final_grade)
                recovered = live.evaluate(1, crop(), evidence, stabilize=True, frame_id=3)
                self.assertEqual(recovered.final_grade, "Class A")
        live = self.engine(minimum_track_observations=2, final_verdict_threshold=.25)
        for frame_id in (1, 2):
            no_body = live.evaluate(2, crop(), [part("Head", "Class A", .99), part("Tail", "Class A", .99)], stabilize=True, frame_id=frame_id)
        self.assertIsNone(no_body.final_grade)
        self.assertIn("UG_TEMPORAL_EVIDENCE", no_body.reason_codes)
        self.assertEqual(no_body.temporal_stability["evidence_bearing_frame_count"], 0)

    def test_transient_uncertainty_recovers_only_after_new_clear_evidence(self) -> None:
        evidence = [part("Body", "Class A", .9), part("Head", "Class A", .9)]
        for flag in ("association_ambiguous", "parent_uncertain"):
            with self.subTest(flag=flag):
                live = self.engine(minimum_track_observations=2)
                live.evaluate(7, crop(), evidence, stabilize=True, frame_id=1)
                uncertain = live.evaluate(7, crop(), evidence, stabilize=True, frame_id=2, frame_quality={flag: True})
                self.assertIsNone(uncertain.final_grade)
                self.assertIsNone(live.evaluate(7, crop(), evidence, stabilize=True, frame_id=3).final_grade)
                recovered = live.evaluate(7, crop(), evidence, stabilize=True, frame_id=4)
                self.assertEqual(recovered.final_grade, "Class A")
                self.assertGreaterEqual(recovered.temporal_stability["historical_uncertainty_counts"][flag], 1)
                live.discard_except(set())
                self.assertIsNone(live.evaluate(7, crop(), evidence, stabilize=True, frame_id=5).final_grade)

    def test_low_confidence_and_strict_partial_frames_do_not_supply_temporal_evidence(self) -> None:
        live = self.engine(minimum_track_observations=2, final_verdict_threshold=.25)
        weak = [part("Body", "Class A", .20), part("Head", "Class A", .95)]
        for frame_id in (1, 2):
            result = live.evaluate(9, crop(), weak, stabilize=True, frame_id=frame_id)
        self.assertEqual(result.temporal_stability["evidence_bearing_frame_count"], 0)
        self.assertIsNone(result.final_grade)
        strict = self.engine(grading_mode="strict", minimum_track_observations=2)
        for frame_id in (1, 2):
            result = strict.evaluate(9, crop(), [part("Body", "Class A", .9), part("Head", "Class A", .9)], stabilize=True, frame_id=frame_id)
        self.assertEqual(result.temporal_stability["evidence_bearing_frame_count"], 0)
        self.assertIsNone(result.final_grade)

    def test_old_evidence_expires_during_occlusion_and_duplicate_cannot_refresh_it(self) -> None:
        live = self.engine(minimum_track_observations=2, temporal_evidence_limit=3, final_verdict_threshold=.25)
        evidence = [part("Body", "Class A", .9)]
        live.evaluate(4, crop(), evidence, stabilize=True, frame_id=1)
        live.evaluate(4, crop(), [], stabilize=True, frame_id=2)
        live.evaluate(4, crop(), evidence, stabilize=True, frame_id=1)
        result = live.evaluate(4, crop(), [], stabilize=True, frame_id=3)
        self.assertEqual(result.temporal_stability["evidence_bearing_frame_count"], 1)
        self.assertEqual(result.temporal_stability["duplicate_frame_count"], 1)
        result = live.evaluate(4, crop(), [], stabilize=True, frame_id=4)
        self.assertEqual(result.temporal_stability["evidence_bearing_frame_count"], 0)
        self.assertIsNone(result.final_grade)
        result = live.evaluate(4, crop(), evidence, stabilize=True, frame_id=5)
        self.assertIsNone(result.final_grade)

    def test_temporal_window_must_allow_required_evidence(self) -> None:
        with self.assertRaises(ValueError):
            GradingConfig(minimum_track_observations=3, temporal_evidence_limit=2)


class ExplainableExportTests(unittest.TestCase):
    def test_csv_and_xlsx_keep_numeric_weighted_calculation_and_rule_sheets(self) -> None:
        verdict = WeightedGradingEngine().evaluate(128, crop(), [part("Body", "Class A", .59), part("Head", "Class A", .72), part("Tail", "Class A", .64)], stabilize=False)
        event = {
            "track_id": 128, "fish_label": "Fish #128", "timestamp": "2026-09-17T10:00:00+00:00",
            "final_confidence": .858, "quality": verdict.final_grade, "quality_confidence": verdict.final_score,
            "parts": [item.to_dict() for item in (part("Body", "Class A", .59), part("Head", "Class A", .72), part("Tail", "Class A", .64))],
            "analysis": verdict.to_dict(), "processing_time_ms": 12.5,
        }
        row = inspection_row(event)
        self.assertAlmostEqual(float(row["final_weighted_score"]), 63.9)
        self.assertAlmostEqual(float(row["body_weighted_contribution"]), 29.5)
        self.assertAlmostEqual(float(row["body_class_a_contribution"]), 29.5)
        self.assertIn("model2_detector_threshold", row)
        self.assertIn("part_evidence_threshold", row)
        self.assertEqual(row["reason_codes"], "GR_CONFIDENT")
        csv_rows = list(csv.DictReader(StringIO(make_csv([event], DEFAULT_EXPORT_FIELDS).decode("utf-8-sig"))))
        self.assertEqual(csv_rows[0]["AI Final Grade"], "Class A")
        self.assertEqual(csv_rows[0]["All Verdict Reason Codes"], "GR_CONFIDENT")
        workbook = make_xlsx([event], DEFAULT_EXPORT_FIELDS, {"Total Fish": 1}, verdict.to_dict().get("grading_config", {"part_weights": {"Body": .5, "Head": .3, "Tail": .2}}))
        from openpyxl import load_workbook
        loaded = load_workbook(BytesIO(workbook), data_only=True)
        self.assertEqual(loaded.sheetnames, ["Fish Inspections", "Session Summary", "Feature Summary", "Grade Calculation Rules", "Review Queue"])
        self.assertEqual(loaded["Fish Inspections"].cell(row=2, column=3).value, "Class A")

    def test_history_keeps_an_updated_calculation_even_when_grade_is_unchanged(self) -> None:
        tracker = TrackingManager(TrackingConfig(line_position=.5))
        first = Detection(0, "Fish", .9, (20, 5, 40, 25), 7)
        crossed = Detection(0, "Fish", .9, (55, 5, 75, 25), 7)
        tracker.update([first], (100, 40), timestamp=0)
        tracker.update([crossed], (100, 40), timestamp=.1, grades={7: QualitySummary(7, "Class A", .5, analysis={"final_score": .5})})
        tracker.update([Detection(0, "Fish", .9, (65, 5, 85, 25), 7)], (100, 40), timestamp=.2, grades={7: QualitySummary(7, "Class A", .6, analysis={"final_score": .6})})
        self.assertEqual(tracker.history()[0]["analysis"]["final_score"], .6)


if __name__ == "__main__":
    unittest.main()
